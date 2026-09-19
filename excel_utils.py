from __future__ import annotations

import calendar
import difflib
import io
import re
from copy import copy
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional

import openpyxl
import pandas as pd
from openpyxl.formula.translate import Translator


MONTH_PATTERN = re.compile(
    r"\b"
    r"(january|february|march|april|may|june|july|august|september|"
    r"october|november|december|jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec)"
    r"\s*,?\s*"
    r"(20\d{2})\b",
    re.IGNORECASE,
)


HEADER_SYNONYMS = {
    "employee_id": {
        "employee id",
        "employee code",
        "emp code",
        "emp. code",
        "code",
        "id",
    },
    "name": {
        "employee name",
        "name of employee",
        "employee",
        "name",
        "name of consultant",
        "consultant",
        "consultant name",
        "name of freelancer",
        "freelancer",
    },
    "essl_present_days": {
        "essl present days",
        "essl present",
    },
    "payroll_present_days": {
        "payroll present days",
        "payroll present",
        "payable days",
        "payroll payable days",
    },
    "days_present": {
        "days present",
        "no days present",
        "dayspresent",
        "present days",
        "present",
        "attendance days",
        "working days present",
        "no. days present",
    },
    "absent_count": {
        "absent count",
        "absence count",
        "days absent",
        "absent",
        "a",
        "essl absent count",
    },
    "payroll_paid_dates": {
        "payroll paid dates",
        "paid dates",
        "policy paid dates",
    },
    "payroll_unpaid_absent_days": {
        "payroll unpaid absent days",
        "unpaid absent days",
    },
    "absent_days": {
        "absent days",
        "absence days",
        "essl absent days",
    },
    "payroll_unpaid_absent_dates": {
        "payroll unpaid absent dates",
        "unpaid absent dates",
    },
    "active_from": {
        "active from",
        "active start",
        "join date",
        "joining date",
    },
    "active_to": {
        "active to",
        "active end",
    },
}


LEAVE_HEADER_ALIASES = {
    "name": {
        "name",
        "employee name",
        "employee",
    },
    "approved_leaves": {
        "no of approved leaves",
        "no. of approved leaves",
        "approved leaves",
        "approved leave",
    },
    "approved_comp_offs": {
        "no of approved comp offs taken",
        "no. of approved comp offs taken",
        "approved comp offs taken",
        "approved comp off taken",
        "comp offs taken",
        "comp off taken",
    },
    "unpaid_leaves": {
        "no of unpaid leaves",
        "no. of unpaid leaves",
        "unpaid leaves",
        "unpaid leave",
    },
}


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = text.replace("\xa0", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _find_header_columns(
    row: Iterable[Any],
) -> Dict[str, int]:
    normalized = [
        _normalize_header(value)
        for value in row
    ]

    columns: Dict[str, int] = {}

    for idx, value in enumerate(normalized):
        if not value:
            continue

        for field_name, variants in HEADER_SYNONYMS.items():
            if field_name in columns:
                continue

            normalized_variants = {
                _normalize_header(alias)
                for alias in variants
            }

            if value in normalized_variants:
                columns[field_name] = idx
                break

    return columns


def _find_leave_header_columns(
    row: Iterable[Any],
) -> Dict[str, int]:
    normalized = [
        _normalize_header(value)
        for value in row
    ]

    columns: Dict[str, int] = {}

    for idx, value in enumerate(normalized):
        if not value:
            continue

        for field_name, variants in LEAVE_HEADER_ALIASES.items():
            if field_name in columns:
                continue

            normalized_variants = {
                _normalize_header(alias)
                for alias in variants
            }

            if value in normalized_variants:
                columns[field_name] = idx
                break

            if (
                field_name == "approved_leaves"
                and "approved" in value
                and "leave" in value
            ):
                columns[field_name] = idx
                break

            if (
                field_name == "approved_comp_offs"
                and "approved" in value
                and "comp" in value
                and "off" in value
            ):
                columns[field_name] = idx
                break

            if (
                field_name == "unpaid_leaves"
                and "unpaid" in value
                and "leave" in value
            ):
                columns[field_name] = idx
                break

    return columns


def _extract_month_year_from_value(
    value: Any,
) -> Optional[str]:
    if value is None:
        return None

    match = MONTH_PATTERN.search(
        str(value)
    )

    if not match:
        return None

    return (
        f"{match.group(1).capitalize()} "
        f"{match.group(2)}"
    )


def infer_month_year(
    file_content: Any = None,
    filename: Optional[str] = None,
) -> Optional[str]:
    """
    Infer month/year from filename, workbook cells, or sheet names.
    """

    if filename:
        month = _extract_month_year_from_value(
            filename.replace("_", " ")
        )
        if month:
            return month

    try:
        if isinstance(
            file_content,
            (bytes, bytearray),
        ):
            source = io.BytesIO(
                file_content
            )
        else:
            source = file_content

        if source is None:
            return None

        if hasattr(source, "seek"):
            source.seek(0)

        xls = pd.ExcelFile(source)

        for sheet_name in xls.sheet_names:
            month = _extract_month_year_from_value(
                sheet_name.replace("_", " ")
            )
            if month:
                return month

        for sheet_name in xls.sheet_names:
            df = pd.read_excel(
                xls,
                sheet_name=sheet_name,
                header=None,
                nrows=30,
            )

            for value in df.to_numpy().flatten():
                month = _extract_month_year_from_value(
                    value
                )
                if month:
                    return month

    except Exception:
        pass

    return None


def _coerce_numeric(
    value: Any,
) -> float:
    if value is None:
        return 0.0

    try:
        if pd.isna(value):
            return 0.0
    except Exception:
        pass

    try:
        return float(value)
    except (
        TypeError,
        ValueError,
    ):
        return 0.0


def _normalize_employee_id_local(
    value: Any,
) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip().replace(",", "")

    if not text:
        return ""

    try:
        number = float(text)

        if number.is_integer():
            return str(int(number))
    except (
        TypeError,
        ValueError,
    ):
        pass

    return text


def _read_dataframe_source(
    source: Any,
) -> pd.DataFrame:
    if isinstance(source, pd.DataFrame):
        return source.copy()

    if hasattr(source, "seek"):
        try:
            source.seek(0)
        except Exception:
            pass

    if isinstance(
        source,
        (bytes, bytearray),
    ):
        source = io.BytesIO(source)

    return pd.read_excel(
        source,
        header=None,
    )


def _valid_attendance_row(
    employee_id: str,
    name: str,
) -> bool:
    """
    Reject obvious ESSL system/default rows.

    A genuine employee may have a numeric ESSL ID, but a name that is
    entirely numeric is treated as a malformed/system row.
    """

    if not employee_id or not name:
        return False

    normalized_name = name.strip()

    if normalized_name.isdigit():
        return False

    if normalized_name.lower() in {
        "total",
        "department default",
    }:
        return False

    return True


def _parse_day_column_header(
    value: Any,
) -> Optional[int]:
    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        try:
            day = int(value)

            if 1 <= day <= 31:
                return day
        except Exception:
            return None

    text = str(value).strip()

    match = re.fullmatch(
        r"(?:day\s*)?(\d{1,2})",
        text,
        flags=re.IGNORECASE,
    )

    if not match:
        return None

    day = int(match.group(1))

    if 1 <= day <= 31:
        return day

    return None


def parse_attendance_input(
    file_content: Any,
    filename: Optional[str] = None,
):
    """
    Parse the attendance Excel produced by the ESSL Extractor.

    Preferred new format:

        Employee ID
        Employee Name
        ESSL Present Days
        Payroll Present Days
        ...
        Day 1 ... Day 31

    Backward compatibility:
        Days Present / Present Days is accepted when
        Payroll Present Days is absent.
    """

    month_year = infer_month_year(
        file_content=file_content,
        filename=filename,
    )

    if hasattr(file_content, "seek"):
        file_content.seek(0)

    if isinstance(
        file_content,
        (bytes, bytearray),
    ):
        file_content = io.BytesIO(
            file_content
        )

    df = pd.read_excel(
        file_content,
        header=None,
    )

    if df.empty:
        raise ValueError(
            "Attendance workbook is empty."
        )

    header_row_idx = None
    header_columns: Dict[str, int] = {}

    for row_idx in range(
        min(len(df), 40)
    ):
        columns = _find_header_columns(
            df.iloc[row_idx].tolist()
        )

        if (
            "name" in columns
            and (
                "employee_id" in columns
                or "payroll_present_days" in columns
                or "days_present" in columns
            )
        ):
            header_row_idx = row_idx
            header_columns = columns
            break

    if header_row_idx is None:
        raise ValueError(
            "Could not identify the attendance header row."
        )

    name_col = header_columns.get(
        "name"
    )

    id_col = header_columns.get(
        "employee_id"
    )

    essl_present_col = header_columns.get(
        "essl_present_days"
    )

    payroll_present_col = header_columns.get(
        "payroll_present_days"
    )

    days_present_col = header_columns.get(
        "days_present"
    )

    absent_count_col = header_columns.get(
        "absent_count"
    )

    payroll_unpaid_absent_col = (
        header_columns.get(
            "payroll_unpaid_absent_days"
        )
    )

    absent_days_col = header_columns.get(
        "absent_days"
    )

    payroll_unpaid_absent_dates_col = (
        header_columns.get(
            "payroll_unpaid_absent_dates"
        )
    )
    payroll_paid_dates_col = (
        header_columns.get(
            "payroll_paid_dates"
        )
    )

    active_from_col = header_columns.get(
        "active_from"
    )

    active_to_col = header_columns.get(
        "active_to"
    )

    if name_col is None:
        raise ValueError(
            "Attendance workbook does not contain an employee name column."
        )

    header_values = (
        df.iloc[header_row_idx]
        .tolist()
    )

    daily_columns: Dict[int, int] = {}

    for col_idx, value in enumerate(
        header_values
    ):
        day = _parse_day_column_header(
            value
        )

        if day is not None:
            daily_columns[day] = col_idx

    records: List[Dict[str, Any]] = []

    for row_idx in range(
        header_row_idx + 1,
        len(df),
    ):
        row = df.iloc[row_idx]

        if name_col >= len(row):
            continue

        raw_name = row.iloc[name_col]

        if pd.isna(raw_name):
            continue

        name = str(raw_name).strip()

        if not name:
            continue

        if name.upper() in {
            "TOTAL",
            "DEPARTMENT DEFAULT",
        }:
            continue

        if id_col is not None:
            raw_id = (
                row.iloc[id_col]
                if id_col < len(row)
                else None
            )
            employee_id = (
                _normalize_employee_id_local(
                    raw_id
                )
            )
        else:
            employee_id = ""

        if not employee_id and len(row) > 1:
            employee_id = (
                _normalize_employee_id_local(
                    row.iloc[1]
                )
            )

        if not _valid_attendance_row(
            employee_id,
            name,
        ):
            continue

        essl_present_days = None

        if (
            essl_present_col is not None
            and essl_present_col < len(row)
        ):
            value = row.iloc[
                essl_present_col
            ]

            if pd.notna(value):
                essl_present_days = (
                    _coerce_numeric(value)
                )

        payroll_present_days = None

        if (
            payroll_present_col is not None
            and payroll_present_col < len(row)
        ):
            value = row.iloc[
                payroll_present_col
            ]

            if pd.notna(value):
                payroll_present_days = (
                    _coerce_numeric(value)
                )

        # Backward compatibility with old output.
        if payroll_present_days is None:
            if days_present_col is not None and days_present_col < len(row):
                value = row.iloc[
                    days_present_col
                ]

                if pd.notna(value):
                    payroll_present_days = (
                        _coerce_numeric(value)
                    )

        if payroll_present_days is None:
            raise ValueError(
                f"Attendance row for '{name}' has no Payroll Present Days."
            )

        if essl_present_days is None:
            essl_present_days = (
                payroll_present_days
            )

        absent_count = 0.0

        if (
            absent_count_col is not None
            and absent_count_col < len(row)
        ):
            absent_count = _coerce_numeric(
                row.iloc[absent_count_col]
            )

        payroll_unpaid_absent_days = None

        if (
            payroll_unpaid_absent_col is not None
            and payroll_unpaid_absent_col < len(row)
        ):
            value = row.iloc[
                payroll_unpaid_absent_col
            ]

            if pd.notna(value):
                payroll_unpaid_absent_days = (
                    _coerce_numeric(value)
                )

        if payroll_unpaid_absent_days is None:
            payroll_unpaid_absent_days = (
                absent_count
            )

        absent_days = ""

        if (
            absent_days_col is not None
            and absent_days_col < len(row)
        ):
            value = row.iloc[
                absent_days_col
            ]

            if pd.notna(value):
                absent_days = str(
                    value
                ).strip()

        payroll_unpaid_absent_dates = ""

        if (
            payroll_unpaid_absent_dates_col is not None
            and payroll_unpaid_absent_dates_col < len(row)
        ):
            value = row.iloc[
                payroll_unpaid_absent_dates_col
            ]

            if pd.notna(value):
                payroll_unpaid_absent_dates = (
                    str(value).strip()
                )
                        payroll_paid_dates = []

        if (
            payroll_paid_dates_col is not None
            and payroll_paid_dates_col < len(row)
        ):
            value = row.iloc[
                payroll_paid_dates_col
            ]

            if pd.notna(value):
                raw_paid_dates = str(
                    value
                ).strip()

                for item in re.split(
                    r"[,;\n]+",
                    raw_paid_dates,
                ):
                    item = item.strip()

                    if not item:
                        continue

                    parsed_paid_date = pd.to_datetime(
                        item,
                        errors="coerce",
                    )

                    if pd.notna(
                        parsed_paid_date
                    ):
                        payroll_paid_dates.append(
                            pd.Timestamp(
                                parsed_paid_date
                            ).strftime(
                                "%Y-%m-%d"
                            )
                        )

        active_from = None
        active_to = None

        if (
            active_from_col is not None
            and active_from_col < len(row)
        ):
            value = row.iloc[
                active_from_col
            ]

            if pd.notna(value):
                active_from = (
                    _coerce_numeric(value)
                )

        if (
            active_to_col is not None
            and active_to_col < len(row)
        ):
            value = row.iloc[
                active_to_col
            ]

            if pd.notna(value):
                active_to = (
                    _coerce_numeric(value)
                )

        present_dates: List[str] = []
        half_present_dates: List[str] = []

        month_number = None
        year_number = None

        if month_year:
            match = MONTH_PATTERN.search(
                str(month_year)
            )

            if match:
                month_text = match.group(1).lower()
                month_aliases = {
                    "jan": "january",
                    "feb": "february",
                    "mar": "march",
                    "apr": "april",
                    "may": "may",
                    "jun": "june",
                    "jul": "july",
                    "aug": "august",
                    "sep": "september",
                    "sept": "september",
                    "oct": "october",
                    "nov": "november",
                    "dec": "december",
                }
                month_name = month_aliases.get(
                    month_text,
                    month_text,
                )
                month_number = list(
                    calendar.month_name
                ).index(
                    month_name.capitalize()
                )
                year_number = int(
                    match.group(2)
                )

        for day, col_idx in sorted(
            daily_columns.items()
        ):
            if col_idx >= len(row):
                continue

            raw_status = row.iloc[
                col_idx
            ]

            if pd.isna(raw_status):
                continue

            status = str(
                raw_status
            ).strip().upper()

            date_key = str(day)

            if (
                month_number is not None
                and year_number is not None
            ):
                date_key = (
                    f"{year_number:04d}-"
                    f"{month_number:02d}-"
                    f"{day:02d}"
                )

            if status in {
                "P",
                "PRESENT",
            }:
                present_dates.append(
                    date_key
                )

            elif status in {
                "½P",
                "1/2P",
                "0.5P",
                "HALF DAY",
                "HALF",
            }:
                half_present_dates.append(
                    date_key
                )

        records.append(
            {
                "employee_id": employee_id,
                "name": name,
                "days_present": float(
                    payroll_present_days
                ),
                "essl_present_days": float(
                    essl_present_days
                ),
                "payroll_present_days": float(
                    payroll_present_days
                ),
                "absent_count": float(
                    absent_count
                ),
                "payroll_unpaid_absent_days": float(
                    payroll_unpaid_absent_days
                ),
                "payroll_paid_dates": payroll_paid_dates,
                "absent_days": absent_days,
                "payroll_unpaid_absent_dates": (
                    payroll_unpaid_absent_dates
                ),
                "active_from": active_from,
                "active_to": active_to,
                "present_dates": present_dates,
                "half_present_dates": half_present_dates,
            }
        )

    if not records:
        raise ValueError(
            "No valid attendance records were found."
        )

    return month_year, records
def parse_daily_punch_report(
    file_content: Any,
    filename: Optional[str] = None,
):
    """
    Parse ESSL Daily Attendance Report / Summary Report.

    The report contains repeated employee blocks such as:

        Employee Code: 12
        Employee Name: Prince Chande

        Date        InTime   OutTime   Status
        05-Aug-2026 11:21             Absent

    Business rule:
        If either InTime OR OutTime exists for an employee/date,
        that date is considered PRESENT.

    Returns:
        month_year,
        {
            employee_id: {
                "name": employee_name,
                "punch_dates": set(...),
                "punch_details": {
                    date: {
                        "in_time": ...,
                        "out_time": ...,
                    }
                },
            }
        }
    """

    if hasattr(file_content, "seek"):
        file_content.seek(0)

    if isinstance(
        file_content,
        (bytes, bytearray),
    ):
        file_content = io.BytesIO(
            file_content
        )

    try:
        xls = pd.ExcelFile(
            file_content
        )
    except Exception as exc:
        raise ValueError(
            "Could not read the Daily Attendance Report. "
            "For .xls files, make sure xlrd>=2.0.1 is installed."
        ) from exc

    employees: Dict[str, Dict[str, Any]] = {}
    detected_dates: List[pd.Timestamp] = []

    for sheet_name in xls.sheet_names:
        df = pd.read_excel(
            xls,
            sheet_name=sheet_name,
            header=None,
        )

        if df.empty:
            continue

        current_employee_id = ""
        current_employee_name = ""

        for row_idx in range(
            len(df)
        ):
            row = df.iloc[row_idx]

            values = [
                value
                for value in row.tolist()
                if pd.notna(value)
                and str(value).strip()
            ]

            if not values:
                continue

            normalized_values = [
                _normalize_header(value)
                for value in values
            ]

            # --------------------------------------------------------
            # Detect employee header row
            # --------------------------------------------------------

            employee_code_index = None
            employee_name_index = None

            for index, value in enumerate(
                normalized_values
            ):
                if value == "employee code":
                    employee_code_index = index

                elif value == "employee name":
                    employee_name_index = index

            if employee_code_index is not None:
                code_value = None

                if (
                    employee_code_index + 1
                    < len(values)
                ):
                    code_value = values[
                        employee_code_index + 1
                    ]

                employee_id = (
                    _normalize_employee_id_local(
                        code_value
                    )
                )

                name_value = None

                if employee_name_index is not None:
                    if (
                        employee_name_index + 1
                        < len(values)
                    ):
                        name_value = values[
                            employee_name_index + 1
                        ]

                employee_name = (
                    str(name_value).strip()
                    if name_value is not None
                    else ""
                )

                if employee_id:
                    current_employee_id = (
                        employee_id
                    )

                    current_employee_name = (
                        employee_name
                    )

                    employees.setdefault(
                        employee_id,
                        {
                            "name": employee_name,
                            "punch_dates": set(),
                            "punch_details": {},
                        },
                    )

                    if employee_name:
                        employees[
                            employee_id
                        ]["name"] = employee_name

                continue

            # --------------------------------------------------------
            # Ignore rows before an employee block
            # --------------------------------------------------------

            if not current_employee_id:
                continue

            # --------------------------------------------------------
            # The Daily Attendance Report has:
            #
            # col 0 = Date
            # col 2 = InTime
            # col 3 = OutTime
            #
            # We deliberately use positional detection because the
            # report repeats these headers for every employee.
            # --------------------------------------------------------

            if len(row) < 4:
                continue

            raw_date = row.iloc[0]

            parsed_date = pd.to_datetime(
                raw_date,
                errors="coerce",
            )

            if pd.isna(parsed_date):
                continue

            parsed_date = pd.Timestamp(
                parsed_date
            ).normalize()

            in_time = row.iloc[2]
            out_time = row.iloc[3]

            has_in_time = (
                pd.notna(in_time)
                and str(in_time).strip() != ""
            )

            has_out_time = (
                pd.notna(out_time)
                and str(out_time).strip() != ""
            )

            # --------------------------------------------------------
            # Critical rule:
            #
            # InTime OR OutTime = PRESENT
            # --------------------------------------------------------

            if not (
                has_in_time
                or has_out_time
            ):
                continue

            date_key = parsed_date.strftime(
                "%Y-%m-%d"
            )

            employees[
                current_employee_id
            ]["punch_dates"].add(
                date_key
            )

            employees[
                current_employee_id
            ]["punch_details"][
                date_key
            ] = {
                "in_time": (
                    str(in_time).strip()
                    if has_in_time
                    else ""
                ),
                "out_time": (
                    str(out_time).strip()
                    if has_out_time
                    else ""
                ),
                "source_sheet": sheet_name,
            }

            detected_dates.append(
                parsed_date
            )

    if not employees:
        raise ValueError(
            "No employee punch records were found "
            "in the Daily Attendance Report."
        )

    month_year = None

    if detected_dates:
        first_date = min(
            detected_dates
        )

        month_year = (
            f"{first_date.strftime('%B')} "
            f"{first_date.year}"
        )

    return (
        month_year,
        employees,
    )


def apply_daily_punch_corrections(
    attendance_records: List[Dict[str, Any]],
    punch_records: Dict[str, Dict[str, Any]],
):
    """
    Correct ESSL attendance using the Daily Attendance Report.

    Only positive corrections are made.

    If Basic ESSL says:
        A / blank / half-day

    but the Daily Attendance Report contains:
        InTime OR OutTime

    then that date becomes PRESENT.

    Existing PRESENT dates are never counted twice.

    Dates already included in the ESSL Payroll Present Days
    through the extractor's payroll-paid-date policy are also
    not counted twice.

    Returns a summary dictionary.
    """

    corrected_employees = 0
    corrected_days = 0
    already_present = 0
    unmatched_punch_ids = 0
    correction_details = []

    for record in attendance_records:
        employee_id = _normalize_employee_id_local(
            record.get("employee_id")
        )

        if not employee_id:
            continue

        punch_record = punch_records.get(
            employee_id
        )

        if punch_record is None:
            unmatched_punch_ids += 1
            record[
                "punch_correction_count"
            ] = 0
            record[
                "punch_correction_dates"
            ] = []
            continue

        existing_present_dates = set(
            record.get(
                "present_dates",
                [],
            )
            or []
        )

        existing_half_dates = set(
            record.get(
                "half_present_dates",
                [],
            )
            or []
        )

        # These dates are already included in the payroll
        # attendance calculation by the ESSL Extractor.
        payroll_paid_dates = set(
            record.get(
                "payroll_paid_dates",
                [],
            )
            or []
        )

        employee_corrections = []

        for date_key in sorted(
            punch_record.get(
                "punch_dates",
                set(),
            )
        ):

            # Already a full PRESENT date.
            if date_key in existing_present_dates:
                already_present += 1
                continue

            # Already included through the payroll calendar
            # policy, so adding it again would double count.
            if date_key in payroll_paid_dates:
                already_present += 1
                continue

            # Convert half-day to full present.
            if date_key in existing_half_dates:
                delta = 0.5
                existing_half_dates.remove(
                    date_key
                )
            else:
                delta = 1.0

            existing_present_dates.add(
                date_key
            )

            detail = (
                punch_record.get(
                    "punch_details",
                    {},
                ).get(
                    date_key,
                    {},
                )
            )

            employee_corrections.append(
                {
                    "date": date_key,
                    "in_time": detail.get(
                        "in_time",
                        "",
                    ),
                    "out_time": detail.get(
                        "out_time",
                        "",
                    ),
                    "reason": (
                        "InTime or OutTime exists"
                    ),
                }
            )

            record[
                "essl_present_days"
            ] = float(
                record.get(
                    "essl_present_days",
                    0,
                )
                or 0
            ) + delta

            record[
                "payroll_present_days"
            ] = float(
                record.get(
                    "payroll_present_days",
                    record.get(
                        "days_present",
                        0,
                    ),
                )
                or 0
            ) + delta

            record[
                "days_present"
            ] = record[
                "payroll_present_days"
            ]

            corrected_days += delta

        record[
            "present_dates"
        ] = sorted(
            existing_present_dates
        )

        record[
            "half_present_dates"
        ] = sorted(
            existing_half_dates
        )

        record[
            "punch_correction_count"
        ] = len(
            employee_corrections
        )

        record[
            "punch_correction_dates"
        ] = [
            item["date"]
            for item in employee_corrections
        ]

        record[
            "punch_correction_details"
        ] = employee_corrections

        if employee_corrections:
            corrected_employees += 1

            correction_details.append(
                {
                    "employee_id": employee_id,
                    "name": record.get(
                        "name",
                        "",
                    ),
                    "corrections": employee_corrections,
                }
            )

    return {
        "employees_corrected": corrected_employees,
        "days_corrected": corrected_days,
        "already_present_or_paid": already_present,
        "unmatched_employee_ids": unmatched_punch_ids,
        "details": correction_details,
    }

def _leave_name_normalized(
    value: Any,
) -> str:
    text = str(
        value or ""
    ).strip().lower()
    text = re.sub(
        r"[^a-z0-9\s]",
        " ",
        text,
    )
    text = re.sub(
        r"\s+",
        " ",
        text,
    ).strip()
    return text


def _leave_name_similarity(
    a: str,
    b: str,
) -> float:
    na = _leave_name_normalized(a)
    nb = _leave_name_normalized(b)

    if not na or not nb:
        return 0.0

    if na == nb:
        return 1.0

    ta = na.split()
    tb = nb.split()

    if (
        ta
        and tb
        and ta[0] == tb[0]
        and len(ta) >= 2
        and len(tb) >= 2
    ):
        if (
            len(ta[1]) == 1
            and tb[1].startswith(ta[1])
        ):
            return 0.95

        if (
            len(tb[1]) == 1
            and ta[1].startswith(tb[1])
        ):
            return 0.95

    ratio = difflib.SequenceMatcher(
        None,
        na,
        nb,
    ).ratio()

    token_a = set(ta)
    token_b = set(tb)

    overlap = 0.0

    if token_a and token_b:
        overlap = len(
            token_a & token_b
        ) / len(
            token_a | token_b
        )

    first_name_bonus = (
        0.10
        if ta and tb and ta[0] == tb[0]
        else 0.0
    )

    return min(
        1.0,
        ratio * 0.60
        + overlap * 0.30
        + first_name_bonus,
    )


def parse_leave_report(
    file_content: Any,
    filename: Optional[str] = None,
):
    """
    Parse the current Leave / Comp Off report.

    Exact source format supplied by the user:

        Name
        No. of Approved Leaves
        No. of Approved Comp Offs Taken
        No. of Unpaid Leaves

    The report currently supplies aggregate counts, not leave dates.
    """

    month_year = infer_month_year(
        file_content=file_content,
        filename=filename,
    )

    if hasattr(file_content, "seek"):
        file_content.seek(0)

    if isinstance(
        file_content,
        (bytes, bytearray),
    ):
        file_content = io.BytesIO(
            file_content
        )

    df = pd.read_excel(
        file_content,
        header=None,
    )

    if df.empty:
        raise ValueError(
            "Leave / Comp Off workbook is empty."
        )

    header_row_idx = None
    header_columns: Dict[str, int] = {}

    for row_idx in range(
        min(len(df), 30)
    ):
        columns = _find_leave_header_columns(
            df.iloc[row_idx].tolist()
        )

        if (
            "name" in columns
            and (
                "approved_leaves" in columns
                or "approved_comp_offs" in columns
                or "unpaid_leaves" in columns
            )
        ):
            header_row_idx = row_idx
            header_columns = columns
            break

    if header_row_idx is None:
        raise ValueError(
            "Could not identify the Leave / Comp Off report header."
        )

    name_col = header_columns["name"]

    approved_leaves_col = header_columns.get(
        "approved_leaves"
    )

    approved_comp_offs_col = header_columns.get(
        "approved_comp_offs"
    )

    unpaid_leaves_col = header_columns.get(
        "unpaid_leaves"
    )

    records: List[Dict[str, Any]] = []

    for row_idx in range(
        header_row_idx + 1,
        len(df),
    ):
        row = df.iloc[row_idx]

        if name_col >= len(row):
            continue

        raw_name = row.iloc[
            name_col
        ]

        if pd.isna(raw_name):
            continue

        name = str(
            raw_name
        ).strip()

        if not name:
            continue

        if name.upper() in {
            "TOTAL",
            "GRAND TOTAL",
        }:
            continue

        approved_leaves = (
            _coerce_numeric(
                row.iloc[
                    approved_leaves_col
                ]
            )
            if (
                approved_leaves_col is not None
                and approved_leaves_col < len(row)
            )
            else 0.0
        )

        approved_comp_offs = (
            _coerce_numeric(
                row.iloc[
                    approved_comp_offs_col
                ]
            )
            if (
                approved_comp_offs_col is not None
                and approved_comp_offs_col < len(row)
            )
            else 0.0
        )

        unpaid_leaves = (
            _coerce_numeric(
                row.iloc[
                    unpaid_leaves_col
                ]
            )
            if (
                unpaid_leaves_col is not None
                and unpaid_leaves_col < len(row)
            )
            else 0.0
        )

        records.append(
            {
                "name": name,
                "approved_leaves": approved_leaves,
                "approved_comp_offs": approved_comp_offs,
                "unpaid_leaves": unpaid_leaves,
            }
        )

    if not records:
        raise ValueError(
            "No Leave / Comp Off records were found."
        )

    return month_year, records


def match_leave_record(
    target_name: str,
    leave_records: List[Dict[str, Any]],
):
    """
    Return:

        record, match_method, score

    Fuzzy matches are only accepted when the best candidate is
    sufficiently strong and clearly separated from the second best.
    """

    target = _leave_name_normalized(
        target_name
    )

    if not target:
        return None, "NO NAME", 0.0

    exact = [
        record
        for record in leave_records
        if _leave_name_normalized(
            record.get("name")
        ) == target
    ]

    if len(exact) == 1:
        return exact[0], "EXACT NAME", 1.0

    if len(exact) > 1:
        return (
            None,
            "AMBIGUOUS EXACT NAME",
            1.0,
        )

    candidates = []

    for record in leave_records:
        score = _leave_name_similarity(
            target_name,
            str(
                record.get("name") or ""
            ),
        )

        candidates.append(
            (
                score,
                record,
            )
        )

    candidates.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    if not candidates:
        return None, "NOT FOUND", 0.0

    best_score, best_record = candidates[0]

    second_score = (
        candidates[1][0]
        if len(candidates) > 1
        else 0.0
    )

    if (
        best_score >= 0.90
        and (
            len(candidates) == 1
            or best_score - second_score >= 0.05
        )
    ):
        return (
            best_record,
            "NAME SIMILARITY",
            best_score,
        )

    return None, "REVIEW", best_score


def apply_leave_adjustments(
    mappings: Dict[str, Dict[str, Any]],
    leave_records: List[Dict[str, Any]],
    total_days: int,
    opening_balances: Optional[Dict[str, float]] = None,
    monthly_accrual: float = 2.0,
) -> Dict[str, int]:
    """Apply the Leave / Comp Off second layer to ESSL payroll days.

    ESSL payroll_present_days is already the corrected attendance value and
    already includes paid Sundays, qualifying Saturdays, and public holidays.
    This function therefore adds only paid leave and approved comp off.

    Leave balance:
        opening balance + monthly accrual = available balance
        paid leave = min(approved leave, available balance)
        unpaid leave = approved leave - paid leave
        closing balance = available balance - paid leave

    The Leave report contains aggregate counts rather than leave dates, so
    overlap with dates already paid by the ESSL calendar policy cannot be
    resolved from this report alone.
    """
    opening_balances = opening_balances or {}
    matched = 0
    unmatched = 0
    adjusted = 0

    for mapping in mappings.values():
        attendance_name = str(mapping.get("attendance_name") or "").strip()
        payroll_name = str(mapping.get("payroll_name") or "").strip()
        target_name = attendance_name or payroll_name

        leave_record, match_method, score = match_leave_record(
            target_name, leave_records
        )

        base_payroll_days = float(
            mapping.get("payroll_present_days")
            if mapping.get("payroll_present_days") is not None
            else mapping.get("days_present") or 0
        )
        mapping["payroll_present_days"] = base_payroll_days

        # Reset all reconciliation fields on every run.
        for key in (
            "approved_leaves", "approved_comp_offs", "reported_unpaid_leaves",
            "balance_based_unpaid_leaves", "unpaid_leaves",
            "opening_leave_balance", "available_leave_balance",
            "paid_leave_used", "closing_leave_balance",
            "leave_adjustment", "comp_off_adjustment",
        ):
            mapping[key] = 0.0
        mapping["monthly_leave_accrual"] = float(monthly_accrual)
        mapping["final_payable_days"] = base_payroll_days
        mapping["leave_match_method"] = match_method
        mapping["leave_match_score"] = score
        mapping["leave_report_name"] = ""
        mapping["leave_note"] = ""

        if leave_record is None:
            unmatched += 1
            if match_method == "REVIEW":
                mapping["leave_note"] = (
                    "Possible Leave report name match found, but it was not "
                    "applied automatically."
                )
            elif match_method == "AMBIGUOUS EXACT NAME":
                mapping["leave_note"] = (
                    "Multiple Leave report rows match this employee."
                )
            else:
                mapping["leave_note"] = (
                    "No matching Leave / Comp Off record found. No leave "
                    "adjustment applied."
                )
            mapping["days_present"] = base_payroll_days
            continue

        matched += 1
        approved_leaves = float(leave_record.get("approved_leaves") or 0)
        approved_comp_offs = float(leave_record.get("approved_comp_offs") or 0)
        reported_unpaid = float(leave_record.get("unpaid_leaves") or 0)

        essl_id = _normalize_employee_id_local(mapping.get("essl_id"))
        opening_balance = float(opening_balances.get(essl_id, 0.0))
        available_balance = opening_balance + float(monthly_accrual)

        paid_leave_used = min(approved_leaves, available_balance)
        balance_based_unpaid = max(0.0, approved_leaves - paid_leave_used)
        total_unpaid = max(reported_unpaid, balance_based_unpaid)
        closing_balance = max(0.0, available_balance - paid_leave_used)

        remaining_capacity = max(
            0.0, float(total_days) - base_payroll_days
        )
        applied_leave = min(paid_leave_used, remaining_capacity)
        remaining_capacity -= applied_leave
        applied_comp_off = min(approved_comp_offs, remaining_capacity)

        final_payable_days = min(
            base_payroll_days + applied_leave + applied_comp_off,
            float(total_days),
        )

        mapping.update({
            "approved_leaves": approved_leaves,
            "approved_comp_offs": approved_comp_offs,
            "reported_unpaid_leaves": reported_unpaid,
            "balance_based_unpaid_leaves": balance_based_unpaid,
            "unpaid_leaves": total_unpaid,
            "opening_leave_balance": opening_balance,
            "available_leave_balance": available_balance,
            "paid_leave_used": paid_leave_used,
            "closing_leave_balance": closing_balance,
            "leave_adjustment": applied_leave,
            "comp_off_adjustment": applied_comp_off,
            "final_payable_days": final_payable_days,
            "leave_report_name": str(leave_record.get("name") or ""),
            "days_present": final_payable_days,
        })

        mapping["leave_note"] = (
            "Leave balance reconciliation applied. "
            f"Opening balance={opening_balance:g}, "
            f"monthly accrual={monthly_accrual:g}, "
            f"available={available_balance:g}, "
            f"approved={approved_leaves:g}, "
            f"paid={paid_leave_used:g}, "
            f"unpaid={total_unpaid:g}, "
            f"closing balance={closing_balance:g}."
        )
        if reported_unpaid != balance_based_unpaid:
            mapping["leave_note"] += (
                " Reported unpaid leave differs from the balance-based calculation."
            )
        if applied_leave < paid_leave_used:
            mapping["leave_note"] += (
                " Paid leave was capped by remaining calendar-day capacity."
            )
        if applied_comp_off < approved_comp_offs:
            mapping["leave_note"] += (
                " Comp Off was capped by remaining calendar-day capacity."
            )

        if applied_leave > 0 or applied_comp_off > 0:
            adjusted += 1

    return {
        "matched": matched,
        "unmatched": unmatched,
        "adjusted": adjusted,
        "total_leave_records": len(leave_records),
    }

def detect_workbook_month(
    wb,
) -> Optional[str]:
    """
    Detect an existing payroll month/year from workbook text.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, 10),
            min_col=1,
            max_col=min(ws.max_column, 12),
        ):
            for cell in row:
                month = _extract_month_year_from_value(
                    cell.value
                )
                if month:
                    return month
    return None


def update_active_workbook_month(
    wb,
    new_month_str: str,
    skip_sheets: Optional[
        Iterable[str]
    ] = None,
):
    skip = {
        str(name)
        for name in (
            skip_sheets or []
        )
    }

    for ws in wb.worksheets:
        if ws.title in skip:
            continue

        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(
                    cell.value,
                    str,
                ):
                    continue

                cell.value = MONTH_PATTERN.sub(
                    lambda match: new_month_str,
                    cell.value,
                )


def analyze_template_sheet(
    sheet,
) -> List[Dict[str, Any]]:
    """
    Identify employee and consultant sections in the payroll template.
    """

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    section_headers = []

    for row_index, row in enumerate(
        rows,
        start=1,
    ):
        normalized = [
            _normalize_header(value)
            for value in row
        ]

        employee_header = None
        consultant_header = None

        for index, value in enumerate(
            normalized
        ):
            if value in {
                "name of employee",
                "employee name",
            }:
                employee_header = index

            elif value in {
                "name of consultant",
                "name of freelancer",
            }:
                consultant_header = index

        if employee_header is not None:
            section_headers.append(
                (
                    row_index,
                    "employee",
                    employee_header + 1,
                )
            )
        elif consultant_header is not None:
            section_headers.append(
                (
                    row_index,
                    "consultant",
                    consultant_header + 1,
                )
            )

    sections = []

    for index, (
        header_row,
        section_type,
        name_col,
    ) in enumerate(section_headers):

        next_header_row = (
            section_headers[index + 1][0]
            if index + 1 < len(section_headers)
            else len(rows) + 1
        )

        section = {
            "type": section_type,
            "header_row": header_row,
            "name_col": name_col,
            "start_row": header_row + 1,
            "end_row": next_header_row - 1,
            "rows": [],
        }

        for row_num in range(
            section["start_row"],
            section["end_row"] + 1,
        ):
            row = rows[
                row_num - 1
            ]

            values_upper = [
                str(value).strip().upper()
                for value in row
                if value is not None
            ]

            if any(
                value == "TOTAL"
                or value.startswith("TOTAL ")
                for value in values_upper
            ):
                section["end_row"] = row_num - 1
                break

            if (
                name_col > len(row)
                or row[
                    name_col - 1
                ] is None
            ):
                continue

            name = str(
                row[
                    name_col - 1
                ]
            ).strip()

            if not name:
                continue

            if name.upper() in {
                "TOTAL",
                "NAME OF EMPLOYEE",
                "NAME OF CONSULTANT",
            }:
                continue

            payroll_code = (
                row[0]
                if len(row) >= 1
                else None
            )

            designation = (
                row[2]
                if len(row) >= 3
                else None
            )

            branch = (
                row[3]
                if len(row) >= 4
                else None
            )

            if section_type == "employee":
                total_days = (
                    row[4]
                    if len(row) >= 5
                    else None
                )

                days_present = (
                    row[5]
                    if len(row) >= 6
                    else None
                )

                gross_salary = (
                    row[6]
                    if len(row) >= 7
                    else 0
                )
            else:
                total_days = (
                    row[3]
                    if len(row) >= 4
                    else None
                )

                days_present = (
                    row[4]
                    if len(row) >= 5
                    else None
                )

                gross_salary = (
                    row[5]
                    if len(row) >= 6
                    else 0
                )

            section["rows"].append(
                {
                    "row_num": row_num,
                    "payroll_code": payroll_code,
                    "name": name,
                    "designation": designation,
                    "branch": branch,
                    "gross_salary": gross_salary,
                    "total_days": total_days,
                    "days_present": days_present,
                }
            )

        sections.append(
            section
        )

    return sections


def _copy_row_with_translated_formulas(
    sheet,
    source_row: int,
    target_row: int,
):
    for col_idx in range(
        1,
        sheet.max_column + 1,
    ):
        source = sheet.cell(
            source_row,
            col_idx,
        )

        target = sheet.cell(
            target_row,
            col_idx,
        )

        if (
            isinstance(
                source.value,
                str,
            )
            and source.value.startswith("=")
        ):
            try:
                target.value = (
                    Translator(
                        source.value,
                        origin=source.coordinate,
                    ).translate_formula(
                        target.coordinate
                    )
                )
            except Exception:
                target.value = source.value
        else:
            target.value = source.value

        if source.has_style:
            target.font = copy(
                source.font
            )
            target.fill = copy(
                source.fill
            )
            target.border = copy(
                source.border
            )
            target.alignment = copy(
                source.alignment
            )
            target.number_format = (
                source.number_format
            )
            target.protection = copy(
                source.protection
            )

        if source.hyperlink:
            target._hyperlink = copy(
                source.hyperlink
            )

        if source.comment:
            target.comment = copy(
                source.comment
            )


def _find_last_data_row(
    section: Dict[str, Any],
) -> Optional[int]:
    rows = section.get(
        "rows",
        [],
    )

    if not rows:
        return None

    return max(
        int(row["row_num"])
        for row in rows
    )


def add_new_hire_to_section(
    sheet,
    section: Dict[str, Any],
    hire: Dict[str, Any],
    total_days: int,
    days_present: float,
) -> Optional[int]:
    start_row = int(
        section["start_row"]
    )

    end_row = int(
        section["end_row"]
    )

    name_col = int(
        section["name_col"]
    )

    empty_row = None

    for row_num in range(
        start_row,
        end_row + 1,
    ):
        value = sheet.cell(
            row_num,
            name_col,
        ).value

        if (
            value is None
            or not str(value).strip()
        ):
            empty_row = row_num
            break

    if empty_row is None:
        return None

    source_row = _find_last_data_row(
        section
    )

    if source_row is not None:
        _copy_row_with_translated_formulas(
            sheet,
            source_row,
            empty_row,
        )

    previous_sr = sheet.cell(
        empty_row - 1,
        1,
    ).value

    try:
        sr_no = (
            int(
                float(
                    previous_sr
                )
            )
            + 1
        )
    except (
        TypeError,
        ValueError,
    ):
        sr_no = ""

    sheet.cell(
        empty_row,
        1,
    ).value = sr_no

    sheet.cell(
        empty_row,
        name_col,
    ).value = hire["name"]

    sheet.cell(
        empty_row,
        3,
    ).value = hire.get(
        "designation",
        "",
    )

    sheet.cell(
        empty_row,
        4,
    ).value = hire.get(
        "branch",
        "",
    )

    if section["type"] == "employee":
        sheet.cell(
            empty_row,
            5,
        ).value = total_days

        sheet.cell(
            empty_row,
            6,
        ).value = days_present

        sheet.cell(
            empty_row,
            7,
        ).value = hire.get(
            "gross_salary",
            0,
        )

        sheet.cell(
            empty_row,
            19,
        ).value = None

    else:
        sheet.cell(
            empty_row,
            4,
        ).value = total_days

        sheet.cell(
            empty_row,
            5,
        ).value = days_present

        sheet.cell(
            empty_row,
            6,
        ).value = hire.get(
            "gross_salary",
            0,
        )

        sheet.cell(
            empty_row,
            15,
        ).value = None

    return empty_row


def get_days_in_month(
    month_year_str: str,
) -> int:
    if not month_year_str:
        return 30

    match = MONTH_PATTERN.search(
        str(month_year_str)
    )

    if not match:
        return 30

    month_name = match.group(1).lower()

    year = int(
        match.group(2)
    )

    month_map = {
        "january": 1,
        "february": 2,
        "march": 3,
        "april": 4,
        "may": 5,
        "june": 6,
        "july": 7,
        "august": 8,
        "september": 9,
        "october": 10,
        "november": 11,
        "december": 12,
        "jan": 1,
        "feb": 2,
        "mar": 3,
        "apr": 4,
        "jun": 6,
        "jul": 7,
        "aug": 8,
        "sep": 9,
        "sept": 9,
        "oct": 10,
        "nov": 11,
        "dec": 12,
    }

    month_number = month_map.get(
        month_name
    )

    if month_number is None:
        return 30

    return calendar.monthrange(
        year,
        month_number,
    )[1]
